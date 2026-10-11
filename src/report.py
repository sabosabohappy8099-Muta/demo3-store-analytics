"""月次レポート(Excel)の生成と、所見文用の要約データの作成。"""

from __future__ import annotations

from dataclasses import replace
from datetime import date, datetime, timedelta
from io import BytesIO

import duckdb
import pandas as pd
from openpyxl import Workbook
from openpyxl.chart import LineChart, Reference
from openpyxl.formatting.rule import ColorScaleRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from src.ai_comment import Comment
from src.config import COMPANY_NAME, WEEKDAY_LABELS
from src.queries import (
    QueryFilter,
    comparison_period,
    composition,
    daily_metrics,
    kpi_with_comparison,
    month_period,
    region_table,
)

# ---------------------------------------------------------------------------
# 書式
# ---------------------------------------------------------------------------
FONT_NAME = "游ゴシック"
ACCENT = "B9832A"
ACCENT_LIGHT = "F3E6CC"
SUBTOTAL_FILL = "FAF6EE"
BORDER_COLOR = "D9D2C3"

FMT_SEN = '#,##0,;[Red]-#,##0,'  # 千円単位(値は円のまま、表示だけ千円)
FMT_INT = "#,##0"
FMT_PCT = '+0.0%;[Red]-0.0%;0.0%'
FMT_SHARE = "0.0%"

_thin = Side(style="thin", color=BORDER_COLOR)
BORDER = Border(left=_thin, right=_thin, top=_thin, bottom=_thin)


def _font(size: float = 11, bold: bool = False, color: str = "2B2B2B") -> Font:
    """共通フォント。"""
    return Font(name=FONT_NAME, size=size, bold=bold, color=color)


def _write_table(ws: Worksheet, top: int, left: int, headers: list[str], rows: list[list], formats: list[str | None],
                 subtotal_rows: set[int] | None = None) -> int:
    """見出し付きの表を書き込み、最後の行番号を返す。

    formats は列ごとの数値書式(None は文字列)。subtotal_rows は太字・背景色を付ける行(0始まり)。
    """
    for j, h in enumerate(headers):
        c = ws.cell(row=top, column=left + j, value=h)
        c.font = _font(10.5, bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor=ACCENT)
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        c.border = BORDER
    for i, row in enumerate(rows, start=1):
        is_sub = subtotal_rows is not None and (i - 1) in subtotal_rows
        for j, v in enumerate(row):
            if isinstance(v, float) and pd.isna(v):
                v = None
            c = ws.cell(row=top + i, column=left + j, value=v)
            c.font = _font(10.5, bold=is_sub)
            c.border = BORDER
            if formats[j]:
                c.number_format = formats[j]
                c.alignment = Alignment(horizontal="right")
            if is_sub:
                c.fill = PatternFill("solid", fgColor=SUBTOTAL_FILL)
    ws.row_dimensions[top].height = 30
    return top + len(rows)


def _title(ws: Worksheet, title: str, subtitle: str, width_cols: int) -> None:
    """シートの見出し(1〜2行目)を書く。"""
    ws["A1"] = title
    ws["A1"].font = _font(16, bold=True)
    ws["A2"] = subtitle
    ws["A2"].font = _font(10, color="6B6A65")
    ws.row_dimensions[1].height = 28
    for col in range(1, width_cols + 1):
        ws.cell(row=3, column=col).border = Border(top=Side(style="medium", color=ACCENT))


def _widths(ws: Worksheet, widths: list[float]) -> None:
    """列幅を設定する。"""
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w


def _page_setup(ws: Worksheet, landscape: bool = True) -> None:
    """印刷設定(A4・横幅に合わせて1ページ)。"""
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.orientation = "landscape" if landscape else "portrait"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.print_options.horizontalCentered = True


# ---------------------------------------------------------------------------
# 要約データ(AI に渡すのはこれだけ)
# ---------------------------------------------------------------------------
def _round(v: float | None, digits: int = 4) -> float | None:
    """JSON にしやすいよう丸める(欠損は None)。"""
    return None if v is None or pd.isna(v) else round(float(v), digits)


def build_monthly_summary(con: duckdb.DuckDBPyConnection, month_start: date) -> dict:
    """対象月の主要な集計結果を、所見文の生成に使う要約データ(辞書)にまとめる。

    生データ(日次・会員)は含めず、全社・都道府県・市区町村・店舗上位/下位などの集計値だけを入れる。
    """
    cur = month_period(month_start)
    prev_year = comparison_period(cur, "前年同期間")
    prev_month = comparison_period(cur, "前期間")
    q = QueryFilter()

    kpi = kpi_with_comparison(con, q, cur, prev_year)
    kpi_pm = kpi_with_comparison(con, q, cur, prev_month)
    total = {
        "sales": _round(kpi["sales"]["current"], 0),
        "sales_yoy": _round(kpi["sales"]["growth"]),
        "sales_mom": _round(kpi_pm["sales"]["growth"]),
        "customers": _round(kpi["customers"]["current"], 0),
        "customers_yoy": _round(kpi["customers"]["growth"]),
        "unit_price": _round(kpi["unit_price"]["current"], 0),
        "unit_price_yoy": _round(kpi["unit_price"]["growth"]),
        "stores": kpi["stores"]["current"],
        "sales_per_store": _round(kpi["sales_per_store"]["current"], 0),
    }

    # 地域・業態の比較は、新店の影響を除くため「既存店」(前年同月より前に開店した店舗)で行う
    existing = QueryFilter(store_ids=existing_store_ids(con, prev_year.start))
    same_store = kpi_with_comparison(con, existing, cur, prev_year)
    total["same_store_sales_yoy"] = _round(same_store["sales"]["growth"])
    prefs = region_table(con, existing, cur, prev_year, "prefecture")
    cities = region_table(con, existing, cur, prev_year, "city").dropna(subset=["growth"]).sort_values("growth")
    stores = region_table(con, existing, cur, prev_year, "store").dropna(subset=["growth"]).sort_values("growth")
    formats = _format_table(con, cur, prev_year, existing)

    ch_cur = composition(con, q, cur, "channel").set_index("key")["share"]
    ch_prev = composition(con, q, prev_year, "channel").set_index("key")["share"]

    store_master = con.execute("SELECT store_id, store_name, open_date FROM stores").df()
    new_stores = store_master[
        pd.to_datetime(store_master["open_date"]) > pd.Timestamp(cur.end) - pd.Timedelta(days=365)
    ].sort_values("open_date")
    closed = con.execute(
        """
        WITH d AS (SELECT store_id, store_name, date, sum(sales) AS v FROM sales
                   WHERE date BETWEEN ? AND ? GROUP BY ALL)
        SELECT store_name, count(*) AS days FROM d WHERE v = 0 GROUP BY 1 ORDER BY 1
        """,
        [cur.start, cur.end],
    ).df()

    def _items(df: pd.DataFrame) -> list[dict]:
        return [{"name": r.name, "sales": _round(r.sales, 0), "yoy": _round(r.growth)} for r in df.itertuples()]

    return {
        "company": COMPANY_NAME,
        "month_label": f"{month_start.year}年{month_start.month}月",
        "period": cur.label(),
        "compare_period": prev_year.label(),
        "total": total,
        "comparison_basis": "地域・業態・店舗の前年比は既存店ベース(前年同月より前に開店した店舗のみ)",
        "prefectures": _items(prefs),
        "cities_top": _items(cities.iloc[::-1].head(3)),
        "cities_bottom": _items(cities.head(3)),
        "stores_top": _items(stores.iloc[::-1].head(3)),
        "stores_bottom": _items(stores.head(3)),
        "formats": _items(formats),
        "channel_share": {k: _round(v) for k, v in ch_cur.items()},
        "mobile_share": {
            "current": _round(ch_cur.get("モバイルオーダー")),
            "previous": _round(ch_prev.get("モバイルオーダー")),
        },
        "new_stores": [
            {"name": r.store_name, "open_date": f"{pd.Timestamp(r.open_date):%Y/%m/%d}"} for r in new_stores.itertuples()
        ],
        "closed_stores": [{"name": r.store_name, "days": int(r.days)} for r in closed.itertuples()],
    }


def existing_store_ids(con: duckdb.DuckDBPyConnection, before: date) -> tuple[str, ...]:
    """指定日より前に開店した店舗(既存店)の ID。"""
    rows = con.execute("SELECT store_id FROM stores WHERE open_date < ? ORDER BY 1", [before]).fetchall()
    return tuple(r[0] for r in rows)


def _format_table(con: duckdb.DuckDBPyConnection, cur, cmp, base: QueryFilter | None = None) -> pd.DataFrame:
    """業態別の売上・前年比(region_table と同じ列名)。base で対象店舗を絞れる。"""
    base = base or QueryFilter()
    rows = []
    for fmt in ("駅ナカ型", "ロードサイド型", "商業施設型"):
        k = kpi_with_comparison(con, replace(base, formats=(fmt,)), cur, cmp)
        rows.append({"name": fmt, "sales": k["sales"]["current"], "growth": k["sales"]["growth"],
                     "customers": k["customers"]["current"], "unit_price": k["unit_price"]["current"],
                     "stores": k["stores"]["current"]})
    return pd.DataFrame(rows)


# ---------------------------------------------------------------------------
# Excel 生成
# ---------------------------------------------------------------------------
def build_report(con: duckdb.DuckDBPyConnection, month_start: date, summary: dict, comment: Comment) -> bytes:
    """月次レポートの Excel を作り、バイト列で返す。"""
    cur = month_period(month_start)
    prev_year = comparison_period(cur, "前年同期間")
    prev_month = comparison_period(cur, "前期間")
    month_label = summary["month_label"]
    subtitle = (
        f"{COMPANY_NAME}|対象期間 {cur.label()}(前年同月 {prev_year.label()})|"
        f"作成日時 {datetime.now():%Y/%m/%d %H:%M}|※ダミーデータによるデモ"
    )

    wb = Workbook()
    _sheet_summary(wb.active, con, cur, prev_year, prev_month, month_label, subtitle, summary, comment)
    _sheet_regions(wb.create_sheet(), con, cur, prev_year, month_label, subtitle)
    _sheet_stores(wb.create_sheet(), con, cur, prev_year, month_label, subtitle)
    _sheet_daily(wb.create_sheet(), con, cur, prev_year, month_label, subtitle)

    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _sheet_summary(ws: Worksheet, con, cur, prev_year, prev_month, month_label: str, subtitle: str,
                   summary: dict, comment: Comment) -> None:
    """シート1:サマリー(主要指標・業態別・チャネル別・所見コメント)。"""
    ws.title = "サマリー"
    _title(ws, f"月次実績レポート {month_label}", subtitle, 8)
    q = QueryFilter()
    yoy = kpi_with_comparison(con, q, cur, prev_year)
    mom = kpi_with_comparison(con, q, cur, prev_month)
    existing = QueryFilter(store_ids=existing_store_ids(con, prev_year.start))
    same = kpi_with_comparison(con, existing, cur, prev_year)
    same_mom = kpi_with_comparison(con, existing, cur, prev_month)

    ws["A4"] = "主要指標"
    ws["A4"].font = _font(12, bold=True)
    mobile = summary["mobile_share"]
    rows = [
        ["売上(千円)", yoy["sales"]["current"], yoy["sales"]["compare"], yoy["sales"]["growth"],
         mom["sales"]["compare"], mom["sales"]["growth"]],
        ["客数(人)", yoy["customers"]["current"], yoy["customers"]["compare"], yoy["customers"]["growth"],
         mom["customers"]["compare"], mom["customers"]["growth"]],
        ["客単価(円)", yoy["unit_price"]["current"], yoy["unit_price"]["compare"], yoy["unit_price"]["growth"],
         mom["unit_price"]["compare"], mom["unit_price"]["growth"]],
        ["営業店舗数(店)", yoy["stores"]["current"], yoy["stores"]["compare"], None, mom["stores"]["compare"], None],
        ["店舗あたり売上(千円)", yoy["sales_per_store"]["current"], yoy["sales_per_store"]["compare"],
         yoy["sales_per_store"]["growth"], mom["sales_per_store"]["compare"], mom["sales_per_store"]["growth"]],
        [f"既存店売上(千円・{len(existing.store_ids)}店)", same["sales"]["current"], same["sales"]["compare"],
         same["sales"]["growth"], same_mom["sales"]["compare"], same_mom["sales"]["growth"]],
    ]
    last = _write_table(ws, 5, 1, ["指標", "当月", "前年同月", "前年比", "前月", "前月比"], rows,
                        [None, FMT_INT, FMT_INT, FMT_PCT, FMT_INT, FMT_PCT])
    # 金額の行は千円表示にする
    for r in (6, 10, 11):
        for col in (2, 3, 5):
            ws.cell(row=r, column=col).number_format = FMT_SEN

    # モバイルオーダー比率
    r = last + 1
    for col, (v, fmt) in enumerate(
        [("モバイルオーダー比率", None), (mobile["current"], FMT_SHARE), (mobile["previous"], FMT_SHARE)], start=1
    ):
        c = ws.cell(row=r, column=col, value=v)
        c.font = _font(10.5)
        c.border = BORDER
        if fmt:
            c.number_format = fmt
            c.alignment = Alignment(horizontal="right")
    if mobile["current"] is not None and mobile["previous"] is not None:
        c = ws.cell(row=r, column=4, value=f"{(mobile['current'] - mobile['previous']) * 100:+.1f}pt")
        c.font = _font(10.5)
        c.alignment = Alignment(horizontal="right")
        c.border = BORDER

    # 業態別
    r += 2
    ws.cell(row=r, column=1, value="業態別").font = _font(12, bold=True)
    fmt_df = _format_table(con, cur, prev_year)
    fmt_same = _format_table(con, cur, prev_year, existing).set_index("name")["growth"]
    total_sales = fmt_df["sales"].sum()
    rows = [[f.name, f.stores, f.sales, f.growth, fmt_same.get(f.name), f.customers, f.unit_price,
             f.sales / total_sales if total_sales else None]
            for f in fmt_df.itertuples()]
    r = _write_table(ws, r + 1, 1, ["業態", "店舗数", "売上(千円)", "前年比(全店)", "前年比(既存店)", "客数(人)", "客単価(円)", "構成比"],
                     rows, [None, FMT_INT, FMT_SEN, FMT_PCT, FMT_PCT, FMT_INT, FMT_INT, FMT_SHARE])

    # チャネル別
    r += 2
    ws.cell(row=r, column=1, value="チャネル別").font = _font(12, bold=True)
    ch = composition(con, QueryFilter(), cur, "channel")
    ch_prev = composition(con, QueryFilter(), prev_year, "channel").set_index("key")
    rows = [[c.key, c.sales, (c.sales / ch_prev.loc[c.key, "sales"] - 1) if c.key in ch_prev.index else None, c.share]
            for c in ch.itertuples()]
    r = _write_table(ws, r + 1, 1, ["チャネル", "売上(千円)", "前年比", "構成比"], rows, [None, FMT_SEN, FMT_PCT, FMT_SHARE])

    # 所見コメント
    r += 2
    label = {"MOCK": "ルールベース", "CLAUDE": "Claude", "GEMINI": "Gemini"}[comment.provider]
    ws.cell(row=r, column=1, value=f"所見コメント(自動生成:{label})").font = _font(12, bold=True)
    lines = [comment.headline, *[f"・{p}" for p in comment.points]]
    if comment.note:
        lines.append(f"※{comment.note}")
    for i, text in enumerate(lines, start=1):
        row = r + i
        ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=8)
        c = ws.cell(row=row, column=1, value=text)
        c.font = _font(11, bold=(i == 1), color="6B6A65" if text.startswith("※") else "2B2B2B")
        c.alignment = Alignment(wrap_text=True, vertical="top")
        # 結合セルは自動で高さが変わらないため、文字数から行の高さを見積もる
        ws.row_dimensions[row].height = max(20, 17 * (len(text) // 62 + 1))

    _widths(ws, [26, 15, 15, 13, 15, 12, 12, 10])
    _page_setup(ws, landscape=False)


def _sheet_regions(ws: Worksheet, con, cur, prev_year, month_label: str, subtitle: str) -> None:
    """シート2:都道府県・市区町村別の実績(都道府県ごとに小計行)。"""
    ws.title = "地域別"
    _title(ws, f"都道府県・市区町村別の実績 {month_label}", subtitle, 10)
    city_pref = dict(con.execute("SELECT DISTINCT city, prefecture FROM stores").fetchall())
    prefs = region_table(con, QueryFilter(), cur, prev_year, "prefecture")
    existing = QueryFilter(store_ids=existing_store_ids(con, prev_year.start))
    same_pref = region_table(con, existing, cur, prev_year, "prefecture").set_index("key")["growth"]
    same_city = region_table(con, existing, cur, prev_year, "city").set_index("key")["growth"]
    same_total = kpi_with_comparison(con, existing, cur, prev_year)["sales"]["growth"]
    rows, subtotals = [], set()
    for p in prefs.itertuples():
        subtotals.add(len(rows))
        rows.append([p.key, "(小計)", p.stores, p.sales, p.compare, p.growth, same_pref.get(p.key), p.customers,
                     p.unit_price, p.share])
        cities = region_table(con, QueryFilter(prefectures=(p.key,)), cur, prev_year, "city")
        for c in cities.itertuples():
            share = c.sales / prefs["sales"].sum()
            rows.append([city_pref.get(c.key, p.key), c.key, c.stores, c.sales, c.compare, c.growth,
                         same_city.get(c.key), c.customers, c.unit_price, share])
    total = prefs[["sales", "compare", "customers", "stores"]].sum()
    subtotals.add(len(rows))
    rows.append(["関東合計", "", int(total["stores"]), total["sales"], total["compare"],
                 total["sales"] / total["compare"] - 1 if total["compare"] else None, same_total, total["customers"],
                 total["sales"] / total["customers"] if total["customers"] else None, 1.0])
    last = _write_table(
        ws, 4, 1,
        ["都道府県", "市区町村", "店舗数", "売上(千円)", "前年同月(千円)", "前年比(全店)", "前年比(既存店)", "客数(人)",
         "客単価(円)", "構成比"],
        rows, [None, None, FMT_INT, FMT_SEN, FMT_SEN, FMT_PCT, FMT_PCT, FMT_INT, FMT_INT, FMT_SHARE], subtotals,
    )
    ws.cell(row=last + 2, column=1, value="※既存店=前年同月より前に開店した店舗。新店のみの市区町村は既存店の前年比が空欄になります。").font = _font(9.5, color="6B6A65")
    ws.freeze_panes = "C5"
    _widths(ws, [12, 18, 9, 14, 15, 12, 13, 12, 11, 9])
    _page_setup(ws)


def _sheet_stores(ws: Worksheet, con, cur, prev_year, month_label: str, subtitle: str) -> None:
    """シート3:店舗別の実績(前年比を条件付き書式で色分け)。"""
    ws.title = "店舗別"
    _title(ws, f"店舗別の実績 {month_label}", subtitle, 11)
    df = region_table(con, QueryFilter(), cur, prev_year, "store")
    master = con.execute("SELECT store_id, format, prefecture, city, open_date FROM stores").df().set_index("store_id")
    df = df.join(master, on="key").sort_values("key")
    df["rank"] = df["growth"].rank(ascending=False, method="min")
    rows = [
        [r.key, r.name, r.format, r.prefecture, r.city, r.sales, r.compare, r.growth, r.customers, r.unit_price,
         int(r.rank) if pd.notna(r.rank) else "新店"]
        for r in df.itertuples()
    ]
    last = _write_table(
        ws, 4, 1,
        ["店舗ID", "店舗名", "業態", "都道府県", "市区町村", "売上(千円)", "前年同月(千円)", "前年比", "客数(人)", "客単価(円)", "前年比順位"],
        rows, [None, None, None, None, None, FMT_SEN, FMT_SEN, FMT_PCT, FMT_INT, FMT_INT, FMT_INT],
    )
    # 前年比:-15%(赤)〜0%(白)〜+15%(緑)のグラデーション。符号付きの表示で色だけに頼らない
    ws.conditional_formatting.add(
        f"H5:H{last}",
        ColorScaleRule(start_type="num", start_value=-0.15, start_color="F4B6B6",
                       mid_type="num", mid_value=0, mid_color="FFFFFF",
                       end_type="num", end_value=0.15, end_color="A9DDB5"),
    )
    for row in range(5, last + 1):
        ws.cell(row=row, column=11).alignment = Alignment(horizontal="center")
    ws.auto_filter.ref = f"A4:K{last}"
    ws.freeze_panes = "C5"
    ws.cell(row=last + 2, column=1, value="※前年比の背景色:赤=減少、白=横ばい、緑=増加(±15%で最も濃くなります)。新店は前年同月の実績がないため比較対象外です。").font = _font(9.5, color="6B6A65")
    _widths(ws, [9, 22, 13, 10, 16, 13, 14, 10, 11, 11, 11])
    _page_setup(ws)


def _sheet_daily(ws: Worksheet, con, cur, prev_year, month_label: str, subtitle: str) -> None:
    """シート4:日別推移(前年同日との比較と折れ線グラフ)。"""
    ws.title = "日別推移"
    _title(ws, f"日別の売上推移 {month_label}", subtitle, 8)
    q = QueryFilter()
    d = daily_metrics(con, q, cur)
    p = daily_metrics(con, q, prev_year)
    # 前年同日(日付が存在しない2/29などは欠損)
    prev_map = dict(zip(p["date"], p["sales"]))
    rows = []
    for r in d.itertuples():
        try:
            ly = prev_map.get(r.date.replace(year=r.date.year - 1))
        except ValueError:
            ly = None
        # 曜日の影響をそろえるため、364日前(前年の同じ曜日)の売上も載せる
        lw = con.execute("SELECT sum(sales) FROM sales WHERE date = ?", [r.date - timedelta(days=364)]).fetchone()[0]
        rows.append([r.date, WEEKDAY_LABELS[r.date.weekday()], r.sales, r.customers,
                     r.sales / r.customers if r.customers else None, ly,
                     (r.sales / ly - 1) if ly else None, lw])
    last = _write_table(
        ws, 4, 1, ["日付", "曜日", "売上(円)", "客数(人)", "客単価(円)", "前年同日(円)", "前年同日比", "前年同曜日(円)"],
        rows, ["yyyy/mm/dd", None, FMT_INT, FMT_INT, FMT_INT, FMT_INT, FMT_PCT, FMT_INT],
    )
    for row in range(5, last + 1):
        ws.cell(row=row, column=2).alignment = Alignment(horizontal="center")
        if ws.cell(row=row, column=2).value in ("土", "日"):
            for col in range(1, 9):
                ws.cell(row=row, column=col).fill = PatternFill("solid", fgColor=SUBTOTAL_FILL)
    ws.freeze_panes = "C5"
    _widths(ws, [12, 6, 14, 11, 11, 14, 11, 15])

    chart = LineChart()
    chart.title = "日別売上(当月と前年同曜日)"
    chart.y_axis.title = "売上(円)"
    chart.y_axis.numFmt = "#,##0"
    chart.height, chart.width = 9, 22
    chart.add_data(Reference(ws, min_col=3, min_row=4, max_row=last), titles_from_data=True)
    chart.add_data(Reference(ws, min_col=8, min_row=4, max_row=last), titles_from_data=True)
    chart.set_categories(Reference(ws, min_col=1, min_row=5, max_row=last))
    chart.x_axis.number_format = "m/d"
    chart.series[0].graphicalProperties.line.solidFill = ACCENT
    chart.series[0].graphicalProperties.line.width = 28000
    chart.series[1].graphicalProperties.line.solidFill = "A3A29C"
    chart.series[1].graphicalProperties.line.dashStyle = "dash"
    ws.add_chart(chart, "J4")
    _page_setup(ws)
