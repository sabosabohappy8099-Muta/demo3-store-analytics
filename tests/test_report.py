"""ai_comment.py(所見文)と report.py(要約データ・Excel)のテスト。"""

from __future__ import annotations

import json
from datetime import date
from io import BytesIO

import duckdb
import pytest
from openpyxl import load_workbook

from src import ai_comment
from src.ai_comment import Comment, generate_comment, generate_mock
from src.data_loader import create_connection
from src.report import build_monthly_summary, build_report

MONTH = date(2026, 9, 1)


@pytest.fixture(scope="module")
def con() -> duckdb.DuckDBPyConnection:
    """data/ のダミーデータを読み込んだ接続。"""
    return create_connection()


@pytest.fixture(scope="module")
def summary(con: duckdb.DuckDBPyConnection) -> dict:
    """対象月の要約データ。"""
    return build_monthly_summary(con, MONTH)


class TestSummary:
    def test_is_json_serializable_and_aggregated_only(self, summary: dict) -> None:
        text = json.dumps(summary, ensure_ascii=False)
        # 生データ(会員IDや日次の行)は含めない
        assert "member" not in text and "M0" not in text
        assert len(text) < 6000

    def test_totals_match_queries(self, con: duckdb.DuckDBPyConnection, summary: dict) -> None:
        sales = con.execute("SELECT sum(sales) FROM sales WHERE date BETWEEN '2026-09-01' AND '2026-09-30'").fetchone()[0]
        assert summary["total"]["sales"] == pytest.approx(sales)
        assert summary["month_label"] == "2026年9月"

    def test_existing_store_basis_excludes_new_stores(self, summary: dict) -> None:
        new_names = {s["name"] for s in summary["new_stores"]}
        ranked = {s["name"] for s in summary["stores_top"] + summary["stores_bottom"]}
        assert new_names and not (new_names & ranked)


class TestMockComment:
    def test_contains_key_facts(self, summary: dict) -> None:
        c = generate_mock(summary)
        assert c.provider == "MOCK"
        assert "2026年9月" in c.headline and "前年比" in c.headline
        assert len(c.points) >= 4
        assert any("モバイルオーダー" in p for p in c.points)

    def test_handles_missing_comparison(self, summary: dict) -> None:
        s = json.loads(json.dumps(summary))
        for k in ("sales_yoy", "customers_yoy", "unit_price_yoy", "same_store_sales_yoy"):
            s["total"][k] = None
        s["prefectures"] = [{**p, "yoy": None} for p in s["prefectures"]]
        c = generate_mock(s)
        assert "比較できません" in c.headline


class TestProviderSwitch:
    def test_default_is_mock(self, monkeypatch: pytest.MonkeyPatch, summary: dict) -> None:
        monkeypatch.delenv("AI_PROVIDER", raising=False)
        monkeypatch.setattr(ai_comment, "get_setting", lambda name, default="": default)
        assert generate_comment(summary).provider == "MOCK"

    @pytest.mark.parametrize("provider", ["CLAUDE", "GEMINI"])
    def test_falls_back_to_mock_without_key(self, monkeypatch: pytest.MonkeyPatch, summary: dict, provider: str) -> None:
        monkeypatch.setattr(ai_comment, "get_setting", lambda name, default="": default)
        c = generate_comment(summary, provider)
        assert c.provider == "MOCK" and c.requested == provider
        assert "失敗" in c.note and "API_KEY" in c.note

    def test_falls_back_on_invalid_output(self, monkeypatch: pytest.MonkeyPatch, summary: dict) -> None:
        def broken(_: dict) -> Comment:
            return ai_comment._parse_comment("これはJSONではありません", "CLAUDE")

        monkeypatch.setattr(ai_comment, "generate_claude", broken)
        c = generate_comment(summary, "CLAUDE")
        assert c.provider == "MOCK" and c.note

    def test_uses_api_result_when_successful(self, monkeypatch: pytest.MonkeyPatch, summary: dict) -> None:
        monkeypatch.setattr(
            ai_comment, "generate_gemini",
            lambda _: ai_comment._parse_comment('{"headline": "好調", "points": ["神奈川が牽引"]}', "GEMINI"),
        )
        c = generate_comment(summary, "GEMINI")
        assert c.provider == "GEMINI" and c.points == ["神奈川が牽引"] and not c.note


def test_build_report_sheets(con: duckdb.DuckDBPyConnection, summary: dict) -> None:
    data = build_report(con, MONTH, summary, generate_mock(summary))
    wb = load_workbook(BytesIO(data))
    assert wb.sheetnames == ["サマリー", "地域別", "店舗別", "日別推移"]
    assert wb["サマリー"]["B6"].value == pytest.approx(summary["total"]["sales"])
    stores = wb["店舗別"]
    assert stores.max_row >= 4 + 40
    assert len(stores.conditional_formatting) == 1
    daily = wb["日別推移"]
    assert daily["A5"].value.date() == date(2026, 9, 1)
    assert len(daily._charts) == 1
